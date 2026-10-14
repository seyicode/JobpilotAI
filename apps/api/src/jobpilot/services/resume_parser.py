"""Local resume text extraction and section splitting. No AI or network calls."""

import io
import re
import zipfile
from dataclasses import asdict, dataclass

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_TYPES = ("pdf", "docx", "txt")

# Canonical section key -> headings that introduce it (lowercase, no trailing colon).
SECTION_HEADINGS = {
    "summary": ("summary", "professional summary", "profile", "objective", "about me"),
    "experience": (
        "experience", "work experience", "professional experience", "employment",
        "employment history", "work history", "career history",
    ),
    "education": ("education", "academic background", "qualifications"),
    "skills": ("skills", "technical skills", "core skills", "key skills", "core competencies"),
    "projects": ("projects", "personal projects", "selected projects"),
    "certifications": ("certifications", "certificates", "licenses and certifications"),
    "awards": ("awards", "honors", "honours", "achievements"),
    "publications": ("publications",),
    "volunteering": ("volunteer experience", "volunteering", "volunteer work"),
    "languages": ("languages",),
    "interests": ("interests", "hobbies"),
}
_HEADING_TO_KEY = {h: key for key, headings in SECTION_HEADINGS.items() for h in headings}


class ResumeParseError(ValueError):
    """The file could not be read as a resume. Messages are safe to show users."""


@dataclass(frozen=True)
class Section:
    key: str
    heading: str
    content: str


def detect_file_type(filename: str, data: bytes) -> str:
    """Return pdf, docx, or txt, checking the file contents match the extension."""
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in SUPPORTED_TYPES:
        raise ResumeParseError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if extension == "pdf" and not data.startswith(b"%PDF"):
        raise ResumeParseError("This file is not a valid PDF.")
    if extension == "docx" and not data.startswith(b"PK"):
        raise ResumeParseError("This file is not a valid DOCX document.")
    return extension


def extract_text(file_type: str, data: bytes) -> str:
    if file_type == "pdf":
        text = _extract_pdf(data)
    elif file_type == "docx":
        text = _extract_docx(data)
    else:
        text = _extract_txt(data)
    text = _normalize(text)
    if not text:
        raise ResumeParseError(
            "No text found. Scanned or image-only resumes are not supported yet."
        )
    return text


def split_sections(text: str) -> list[Section]:
    """Split on recognised headings. Text before the first heading becomes 'header'."""
    sections: list[Section] = []
    key, heading, lines = "header", "Header", []
    for line in text.splitlines():
        match = _match_heading(line)
        if match:
            _append(sections, key, heading, lines)
            key, heading, lines = match, line.strip().rstrip(":").strip(), []
        else:
            lines.append(line)
    _append(sections, key, heading, lines)
    return sections


def parse_resume(filename: str, data: bytes) -> tuple[str, str, list[dict]]:
    file_type = detect_file_type(filename, data)
    text = extract_text(file_type, data)
    return file_type, text, [asdict(section) for section in split_sections(text)]


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        encrypted = reader.is_encrypted
        pages = [] if encrypted else [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError, KeyError):
        raise ResumeParseError("This PDF could not be read.") from None
    if encrypted:
        raise ResumeParseError("Password-protected PDFs are not supported.")
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError):
        raise ResumeParseError("This DOCX document could not be read.") from None
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ResumeParseError("This text file uses an unsupported encoding.")


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _match_heading(line: str) -> str | None:
    candidate = re.sub(r"\s+", " ", line.strip().rstrip(":")).lower()
    if not candidate or len(candidate) > 40:
        return None
    return _HEADING_TO_KEY.get(candidate.replace("&", "and"))


def _append(sections: list[Section], key: str, heading: str, lines: list[str]) -> None:
    content = "\n".join(lines).strip()
    if content:
        sections.append(Section(key=key, heading=heading, content=content))
