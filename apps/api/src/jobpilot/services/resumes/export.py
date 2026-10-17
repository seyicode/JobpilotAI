"""ATS-friendly DOCX and PDF rendering of a tailored resume, entirely in memory.

Layout rules for applicant tracking systems: one column, standard section headings, real text
(no images), plain bullets, contact details in the body (no page headers/footers), and no
tables or text boxes. Nothing is written to disk.
"""

import io
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer
from reportlab.platypus.flowables import HRFlowable

from jobpilot.services.resumes.schema import ResumeContent, ResumeEntry

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
PAGE_SIZES = {"A4": A4, "LETTER": LETTER}


def date_range(entry: ResumeEntry) -> str | None:
    def fmt(value: str | None) -> str | None:
        if not value:
            return None
        year, _, month = value.partition("-")
        return f"{MONTHS[int(month) - 1]} {year}" if month else year

    start, end = fmt(entry.start_date), "Present" if entry.is_current else fmt(entry.end_date)
    if start and end:
        return f"{start} – {end}"
    return start or end


def contact_line(content: ResumeContent) -> str:
    h = content.header
    return " | ".join(v for v in (h.email, h.phone, h.location, *h.links) if v)


def entry_heading(entry: ResumeEntry) -> str:
    return f"{entry.title}, {entry.organization}" if entry.organization else entry.title


def entry_meta(entry: ResumeEntry) -> str:
    return " | ".join(v for v in (entry.location, date_range(entry)) if v)


def export_filename(content: ResumeContent, job_title: str | None, company: str | None,
                    version: int, extension: str) -> str:
    parts = [content.header.name or "Resume", job_title, company, f"v{version}"]
    slug = "-".join(re.sub(r"[^A-Za-z0-9]+", "-", p).strip("-") for p in parts if p)
    return f"{slug[:120] or 'Resume'}.{extension}"


# DOCX


def render_docx(content: ResumeContent) -> bytes:
    document = Document()
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(0.75)
        section.top_margin = section.bottom_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    heading = document.styles["Heading 1"]  # Real heading style, so ATS parsers see sections
    heading.font.name = "Calibri"
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(3)

    if content.header.name:
        name = document.add_paragraph()
        run = name.add_run(content.header.name)
        run.bold = True
        run.font.size = Pt(18)
        name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if content.header.headline:
        document.add_paragraph(content.header.headline)
    if contact := contact_line(content):
        document.add_paragraph(contact)

    if content.summary:
        document.add_heading("Summary", level=1)
        for paragraph in content.summary.split("\n\n"):
            document.add_paragraph(paragraph.strip())
    if content.skills:
        document.add_heading("Skills", level=1)
        document.add_paragraph(", ".join(content.skills))
    for section in content.sections:
        document.add_heading(section.heading, level=1)
        for entry in section.entries:
            title = document.add_paragraph()
            title.paragraph_format.space_before = Pt(6)
            title.add_run(entry_heading(entry)).bold = True
            if meta := entry_meta(entry):
                document.add_paragraph().add_run(meta).italic = True
            for bullet in entry.bullets:
                document.add_paragraph(bullet.text, style="List Bullet")

    properties = document.core_properties
    properties.title = f"{content.header.name or 'Resume'} resume"
    properties.author = content.header.name or ""
    properties.comments = ""
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# PDF


def render_pdf(content: ResumeContent, page_size: str = "A4") -> bytes:
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=2)
    styles = {
        "name": ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22),
        "heading": ParagraphStyle("heading", parent=base, fontName="Helvetica-Bold", fontSize=11.5,
                                  leading=14, spaceBefore=10, spaceAfter=2),
        "entry": ParagraphStyle("entry", parent=base, fontName="Helvetica-Bold", spaceBefore=6),
        "meta": ParagraphStyle("meta", parent=base, fontName="Helvetica-Oblique"),
        "bullet": ParagraphStyle("bullet", parent=base, leftIndent=14, bulletIndent=4),
    }

    def text(value: str, style: str, **kwargs) -> Paragraph:
        # Escape so user text is never interpreted as reportlab markup.
        return Paragraph(escape(value), styles.get(style, base), **kwargs)

    def heading(title: str) -> list:
        return [text(title, "heading"), HRFlowable(width="100%", thickness=0.5, spaceAfter=3)]

    story = []
    if content.header.name:
        story.append(text(content.header.name, "name"))
    if content.header.headline:
        story.append(text(content.header.headline, "base"))
    if contact := contact_line(content):
        story.append(text(contact, "base"))
    if content.summary:
        story += heading("Summary")
        story += [text(p.strip(), "base") for p in content.summary.split("\n\n")]
    if content.skills:
        story += heading("Skills")
        story.append(text(", ".join(content.skills), "base"))
    for section in content.sections:
        story += heading(section.heading)
        for entry in section.entries:
            block = [text(entry_heading(entry), "entry")]
            if meta := entry_meta(entry):
                block.append(text(meta, "meta"))
            block += [text(b.text, "bullet", bulletText="•") for b in entry.bullets]
            # Keep an entry's heading with at least its first bullets on one page.
            story.append(KeepTogether(block[:3]))
            story += block[3:]
    if not story:
        story.append(Spacer(1, 1))

    buffer = io.BytesIO()
    margin = 0.7 * inch
    document = SimpleDocTemplate(
        buffer, pagesize=PAGE_SIZES[page_size], leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=margin, title=f"{content.header.name or 'Resume'} resume",
        author=content.header.name or "", creator="JobPilot AI", invariant=1,
    )
    document.build(story)
    return buffer.getvalue()
