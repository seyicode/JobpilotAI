import io
import uuid

import pytest
from docx import Document
from pypdf import PdfReader

from jobpilot.services.resumes.export import (
    date_range,
    export_filename,
    render_docx,
    render_pdf,
)
from jobpilot.services.resumes.schema import ResumeContent, ResumeEntry

CONTENT = ResumeContent.model_validate({
    "header": {"name": "Alex Example", "headline": "Backend Engineer", "email": "alex@example.com",
               "phone": "+44 20 7946 0000", "location": "London, UK", "links": ["https://example.com/alex"]},
    "summary": "Backend engineer focused on payments.\n\nMentor and R&D <lead>.",
    "skills": ["Python", "PostgreSQL", "C++"],
    "sections": [
        {"kind": "job", "heading": "Experience", "entries": [{
            "source_id": str(uuid.uuid4()), "kind": "job", "title": "Senior Engineer",
            "organization": "Example Corp", "location": "London", "start_date": "2021-03", "is_current": True,
            "bullets": [{"text": "Built payment APIs in Python", "source_text": "x"},
                        {"text": "Cut costs by 20% & latency <50ms", "source_text": "x"}],
        }]},
        {"kind": "education", "heading": "Education", "entries": [{
            "source_id": str(uuid.uuid4()), "kind": "education", "title": "BSc Computer Science",
            "organization": "Example University", "end_date": "2016", "bullets": [],
        }]},
    ],
})
READING_ORDER = [
    "Alex Example", "Backend Engineer", "alex@example.com | +44 20 7946 0000 | London, UK | https://example.com/alex",
    "Summary", "Backend engineer focused on payments.", "Mentor and R&D <lead>.",
    "Skills", "Python, PostgreSQL, C++",
    "Experience", "Senior Engineer, Example Corp", "London | Mar 2021 – Present",
    "Built payment APIs in Python", "Cut costs by 20% & latency <50ms",
    "Education", "BSc Computer Science, Example University", "2016",
]


def test_date_ranges():
    entry = ResumeEntry(source_id=uuid.uuid4(), kind="job", title="x", start_date="2019", end_date="2021-12")
    assert date_range(entry) == "2019 – Dec 2021"
    assert date_range(entry.model_copy(update={"start_date": None})) == "Dec 2021"
    assert date_range(entry.model_copy(update={"start_date": None, "end_date": None})) is None


def test_docx_is_single_column_text_in_reading_order():
    document = Document(io.BytesIO(render_docx(CONTENT)))
    paragraphs = [p.text for p in document.paragraphs if p.text]
    assert paragraphs == READING_ORDER
    assert document.tables == []
    assert all(not s.header.paragraphs[0].text and not s.footer.paragraphs[0].text for s in document.sections)
    styles = {p.text: p.style.name for p in document.paragraphs}
    assert styles["Experience"] == "Heading 1"
    assert styles["Built payment APIs in Python"] == "List Bullet"
    assert document.core_properties.author == "Alex Example"


def test_pdf_has_extractable_text_in_reading_order():
    data = render_pdf(CONTENT)
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() for page in reader.pages)
    positions = [text.find(line) for line in READING_ORDER]
    assert -1 not in positions, [line for line, pos in zip(READING_ORDER, positions) if pos == -1]
    assert positions == sorted(positions)
    assert reader.metadata.title == "Alex Example resume"
    assert render_pdf(CONTENT) == data  # deterministic output


def test_pdf_page_sizes():
    a4 = PdfReader(io.BytesIO(render_pdf(CONTENT, "A4"))).pages[0].mediabox
    letter = PdfReader(io.BytesIO(render_pdf(CONTENT, "LETTER"))).pages[0].mediabox
    assert (round(float(a4.width)), round(float(letter.width))) == (595, 612)


def test_empty_resume_still_renders():
    assert PdfReader(io.BytesIO(render_pdf(ResumeContent()))).pages
    assert Document(io.BytesIO(render_docx(ResumeContent())))


def test_filenames_are_safe():
    assert export_filename(CONTENT, "Backend Engineer / Payments", "Acme, Inc.", 3, "pdf") == (
        "Alex-Example-Backend-Engineer-Payments-Acme-Inc-v3.pdf")
    assert export_filename(ResumeContent(), None, None, 1, "docx") == "Resume-v1.docx"


# API


@pytest.fixture
def approved(db_client):
    client = db_client
    client.put("/profile", json={"full_name": "Alex Example"})
    client.post("/experiences", json={"kind": "job", "title": "Engineer", "organization": "Example Corp",
                                      "highlights": ["Built Python APIs"], "verified": True})
    client.post("/skills", json={"name": "Python", "verified": True})
    description = "Backend Engineer at Acme\nRequirements: Python. We build payment systems for merchants."
    job = client.post("/jobs", json={"description": description}).json()["job"]
    version = client.post(f"/jobs/{job['id']}/resumes").json()
    return client, version


def test_drafts_cannot_be_exported(approved):
    client, version = approved
    response = client.get(f"/tailored-resumes/{version['id']}/export?format=pdf")
    assert response.status_code == 409
    assert response.json() == {"detail": "Approve this version before exporting it."}


@pytest.mark.parametrize(("fmt", "media_type"), [
    ("pdf", "application/pdf"),
    ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
])
def test_export_approved_version(approved, fmt, media_type):
    client, version = approved
    client.post(f"/tailored-resumes/{version['id']}/approve")
    response = client.get(f"/tailored-resumes/{version['id']}/export?format={fmt}")
    assert response.status_code == 200
    assert response.headers["content-type"] == media_type
    assert response.headers["cache-control"] == "no-store, private"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="Alex-Example-Backend-Engineer-Acme-v1.{fmt}"')
    assert len(response.content) > 1000
    assert client.get(f"/tailored-resumes/{version['id']}/export?format=txt").status_code == 422
